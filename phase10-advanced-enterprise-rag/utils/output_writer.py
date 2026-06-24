import pandas as pd
from docx import Document


def save_outputs(test_cases, output_folder="outputs"):

    rows = []

    for tc in test_cases:

        rows.append({
            "Test Case ID": tc.test_case_id,
            "Requirement ID": tc.requirement_id,
            "Title": tc.title,
            "Priority": tc.priority,
            "Type": tc.test_type,
            "Preconditions": tc.preconditions,
            "Test Data": tc.test_data,
            "Steps": "\n".join(tc.steps),
            "Expected Result": tc.expected_result,
            "Automation Candidate": tc.automation_candidate,
            "Source Traceability": tc.source_traceability
        })

    df = pd.DataFrame(rows)

    excel_path = f"{output_folder}/test_cases.xlsx"

    df.to_excel(
        excel_path,
        index=False
    )

    word_path = f"{output_folder}/test_cases.docx"

    document = Document()

    document.add_heading(
        "Enterprise Test Cases",
        level=1
    )

    for tc in test_cases:

        document.add_heading(
            f"{tc.test_case_id} - {tc.title}",
            level=2
        )

        document.add_paragraph(f"Requirement ID: {tc.requirement_id}")
        document.add_paragraph(f"Priority: {tc.priority}")
        document.add_paragraph(f"Type: {tc.test_type}")
        document.add_paragraph(f"Preconditions: {tc.preconditions}")
        document.add_paragraph(f"Test Data: {tc.test_data}")

        document.add_paragraph("Steps:")

        for index, step in enumerate(tc.steps, start=1):
            clean_step = str(step).strip()
            document.add_paragraph(f"{index}. {clean_step}")

        document.add_paragraph(f"Expected Result: {tc.expected_result}")
        document.add_paragraph(f"Automation Candidate: {tc.automation_candidate}")
        document.add_paragraph(f"Source Traceability: {tc.source_traceability}")

        document.add_page_break()

    document.save(word_path)

    return {
        "excel": excel_path,
        "word": word_path
    }